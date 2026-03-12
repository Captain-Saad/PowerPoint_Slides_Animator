"""
engine/slide_analyzer.py
Analyzes a python-pptx Presentation: detects the type of each slide
and classifies every shape into a semantic ElementType.

Returns a list of SlideAnalysis objects — one per slide — consumed by
the AnimationPlanner in the next pipeline stage.
"""

from __future__ import annotations

import logging
import sys
from pathlib import Path
from typing import List, Optional, Set

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from models.animation_plan import ShapeInfo, SlideAnalysis
from models.element_type import ElementType
from models.slide_type import SlideType

logger = logging.getLogger(__name__)

# EMU threshold below which a shape is treated as decorative accent.
# 914400 EMU = 1 inch.  Shapes smaller than 2 inches on either side
# with no text are classified as SHAPE_ACCENT (icons, lines, dividers).
_ACCENT_SIZE_THRESHOLD_EMU = 1_828_800  # 2 inches

# Layout name keywords that indicate a section header slide.
_SECTION_KEYWORDS = {"section", "divider", "chapter", "interstitial"}

# PP_PLACEHOLDER types that map directly to TITLE
_TITLE_PH_TYPES: Set[int] = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}

# PP_PLACEHOLDER types that map to BODY text
_BODY_PH_TYPES: Set[int] = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}


class SlideAnalyzer:
    """Analyzes all slides in a python-pptx Presentation object.

    Responsibilities:
    - Open a .pptx file path with python-pptx
    - For each slide: detect its SlideType
    - For each shape on that slide: classify its ElementType
    - Return a list of SlideAnalysis instances (one per slide)

    Usage::

        analyzer = SlideAnalyzer("path/to/deck.pptx")
        analyses = analyzer.analyze_all()
    """

    def __init__(self, pptx_path: str) -> None:
        """Initialize the analyzer with the path to a .pptx file.

        Args:
            pptx_path: Absolute or relative path to the input .pptx file.

        Raises:
            SystemExit: If the file does not exist or cannot be opened as
                a PowerPoint presentation.
        """
        path = Path(pptx_path)
        if not path.exists():
            sys.exit(f"Error: File not found: {pptx_path}")
        try:
            self.prs = Presentation(str(path))
        except Exception as exc:
            sys.exit(
                f"Error: Could not open file as a PowerPoint presentation: "
                f"{pptx_path}\nDetails: {exc}"
            )
        self.pptx_path = path
        logger.debug("Opened presentation: %s", path.name)

    def analyze_all(self) -> List[SlideAnalysis]:
        """Analyze every slide in the presentation.

        Returns:
            Ordered list of :class:`~models.animation_plan.SlideAnalysis`
            objects, one per slide, in presentation order.
        """
        analyses: List[SlideAnalysis] = []
        for idx, slide in enumerate(self.prs.slides):
            analysis = self._analyze_slide(slide, idx)
            analyses.append(analysis)
            logger.debug(
                "Slide %d: type=%s, elements=%d",
                analysis.slide_number,
                analysis.slide_type.value,
                analysis.total_element_count,
            )
        return analyses

    def _analyze_slide(self, slide: object, slide_index: int) -> SlideAnalysis:
        """Analyze a single slide and return its SlideAnalysis.

        Args:
            slide: A python-pptx ``Slide`` object.
            slide_index: Zero-based index of the slide.

        Returns:
            Populated SlideAnalysis for the given slide.
        """
        layout_name: Optional[str] = None
        try:
            layout_name = slide.slide_layout.name
        except Exception:
            pass

        shapes: List[ShapeInfo] = []
        for shape in slide.shapes:
            info = self._classify_shape(shape)
            if info is not None:
                shapes.append(info)

        slide_type = self._detect_slide_type(slide, shapes, layout_name)

        return SlideAnalysis(
            slide_index=slide_index,
            slide_number=slide_index + 1,
            slide_type=slide_type,
            shapes=shapes,
            total_element_count=len(shapes),
            layout_name=layout_name,
        )

    def _detect_slide_type(
        self,
        slide: object,
        shapes: List[ShapeInfo],
        layout_name: Optional[str] = None,
    ) -> SlideType:
        """Determine the semantic SlideType of a slide.

        Detection uses layout name inspection, placeholder types, and the
        presence/absence of text and image elements.

        Priority order:
          1. No shapes → BLANK_SLIDE
          2. Layout name contains a section keyword → SECTION_HEADER
          3. Contains only TITLE + optional SUBTITLE, no BODY/IMAGE → TITLE_SLIDE
          4. All shapes are images only (no text at all) → IMAGE_ONLY
          5. Default → CONTENT_SLIDE

        Args:
            slide: A python-pptx ``Slide`` object.
            shapes: Already-classified shape list for this slide.
            layout_name: The slide layout name string, or None.

        Returns:
            The detected :class:`~models.slide_type.SlideType`.
        """
        if not shapes:
            return SlideType.BLANK_SLIDE

        # Check layout name for section keywords
        if layout_name:
            lower_name = layout_name.lower()
            if any(kw in lower_name for kw in _SECTION_KEYWORDS):
                return SlideType.SECTION_HEADER

        element_types: Set[ElementType] = {s.element_type for s in shapes}

        text_types = {ElementType.TITLE, ElementType.SUBTITLE, ElementType.BODY,
                      ElementType.SHAPE_CONTENT}
        has_title = ElementType.TITLE in element_types
        has_subtitle = ElementType.SUBTITLE in element_types
        has_body = ElementType.BODY in element_types
        has_text_content = bool(element_types & text_types)
        has_image = ElementType.IMAGE in element_types

        # Title + optional subtitle only (no body, no standalone images)
        non_title_text = element_types - {ElementType.TITLE, ElementType.SUBTITLE,
                                           ElementType.SHAPE_ACCENT, ElementType.ICON}
        if has_title and not has_body and not non_title_text - {ElementType.SUBTITLE}:
            return SlideType.TITLE_SLIDE

        # No text shapes at all but has images
        if not has_text_content and has_image:
            return SlideType.IMAGE_ONLY

        # No content to speak of (only accents or unknown)
        meaningful = element_types - {ElementType.SHAPE_ACCENT, ElementType.ICON,
                                       ElementType.UNKNOWN}
        if not meaningful:
            return SlideType.BLANK_SLIDE

        return SlideType.CONTENT_SLIDE

    def _classify_shape(self, shape: object) -> Optional[ShapeInfo]:
        """Classify a single shape into a ShapeInfo.

        Evaluation order (first match wins):
          1. Group shape → GROUPED
          2. Has table → TABLE
          3. Has chart → CHART
          4. Placeholder shapes: type-specific mapping
          5. Non-placeholder PICTURE → IMAGE
          6. Non-placeholder with text → SHAPE_CONTENT
          7. Non-placeholder no text, small → SHAPE_ACCENT

        Args:
            shape: A python-pptx ``BaseShape`` subclass instance.

        Returns:
            A populated :class:`~models.animation_plan.ShapeInfo`, or
            ``None`` if the shape has no usable ``shape_id`` attribute.
        """
        try:
            shape_id: int = shape.shape_id
        except AttributeError:
            logger.warning("Shape has no shape_id — skipping: %s", getattr(shape, "name", "?"))
            return None

        name: str = getattr(shape, "name", f"shape_{shape_id}")
        element_type = self._determine_element_type(shape)
        has_text = self._shape_has_text(shape)
        para_count = self._count_paragraphs(shape)

        return ShapeInfo(
            shape_id=shape_id,
            shape_name=name,
            element_type=element_type,
            paragraph_count=para_count,
            has_text=has_text,
        )

    def _determine_element_type(self, shape: object) -> ElementType:
        """Map a shape to its ElementType using the classification priority order.

        Args:
            shape: A python-pptx shape object.

        Returns:
            The best-matching :class:`~models.element_type.ElementType`.
        """
        # 1. Group
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                return ElementType.GROUPED
        except Exception:
            pass

        # 2. Table (GraphicFrame with table)
        try:
            if shape.has_table:
                return ElementType.TABLE
        except Exception:
            pass

        # 3. Chart (GraphicFrame with chart)
        try:
            if shape.has_chart:
                return ElementType.CHART
        except Exception:
            pass

        # 4. Placeholder shapes
        try:
            if shape.is_placeholder:
                return self._classify_placeholder(shape)
        except Exception:
            pass

        # 5. Non-placeholder picture
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return ElementType.IMAGE
        except Exception:
            pass

        # 6 & 7. Text-bearing vs decorative free shapes
        return self._classify_free_shape(shape)

    def _classify_placeholder(self, shape: object) -> ElementType:
        """Classify a confirmed placeholder shape by its placeholder type.

        Args:
            shape: A python-pptx placeholder shape (``is_placeholder`` is True).

        Returns:
            The appropriate :class:`~models.element_type.ElementType`.
        """
        try:
            ph_type = shape.placeholder_format.type
            if ph_type in _TITLE_PH_TYPES:
                return ElementType.TITLE
            if ph_type == PP_PLACEHOLDER.SUBTITLE:
                return ElementType.SUBTITLE
            if ph_type == PP_PLACEHOLDER.PICTURE:
                return ElementType.IMAGE
            if ph_type in _BODY_PH_TYPES:
                return ElementType.BODY
            # Custom / unknown placeholder — use content heuristic
            if self._shape_has_text(shape):
                return ElementType.BODY
            return ElementType.IMAGE
        except Exception:
            # Can't read placeholder type — fall back
            if self._shape_has_text(shape):
                return ElementType.BODY
            return ElementType.IMAGE

    def _classify_free_shape(self, shape: object) -> ElementType:
        """Classify a non-placeholder, non-group, non-media free shape.

        Shapes with text are SHAPE_CONTENT; shapes without text that are
        small are SHAPE_ACCENT; larger shapes without text are also
        SHAPE_ACCENT (borders, decorative fills, etc.).

        Args:
            shape: A non-placeholder python-pptx shape.

        Returns:
            Either ``SHAPE_CONTENT`` or ``SHAPE_ACCENT``.
        """
        if self._shape_has_text(shape):
            return ElementType.SHAPE_CONTENT

        # Determine size — small shapes without text are accent elements
        try:
            width = shape.width or 0
            height = shape.height or 0
            if width < _ACCENT_SIZE_THRESHOLD_EMU or height < _ACCENT_SIZE_THRESHOLD_EMU:
                return ElementType.SHAPE_ACCENT
        except Exception:
            pass

        return ElementType.SHAPE_ACCENT

    def _shape_has_text(self, shape: object) -> bool:
        """Return True if the shape contains any non-whitespace text.

        Args:
            shape: A python-pptx shape object.

        Returns:
            Boolean indicating whether the shape has real text content.
        """
        try:
            if not shape.has_text_frame:
                return False
            return bool(shape.text_frame.text.strip())
        except Exception:
            return False

    def _count_paragraphs(self, shape: object) -> int:
        """Count the number of non-empty paragraphs in a text frame.

        Used by the AnimationPlanner when ``by_paragraph`` mode is active
        to know how many AnimationEntry objects to create for one shape.

        Args:
            shape: A python-pptx shape that may contain a text frame.

        Returns:
            Number of non-empty paragraphs, or 0 if no text frame.
        """
        try:
            if not shape.has_text_frame:
                return 0
            return sum(
                1 for p in shape.text_frame.paragraphs if p.text.strip()
            )
        except Exception:
            return 0
