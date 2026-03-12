"""
tests/test_slide_analyzer.py
Unit tests for engine/slide_analyzer.py.

All slide fixtures are built programmatically with python-pptx so no
binary .pptx files need to be committed to source control for M2.
"""

from __future__ import annotations

import io
import os
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from engine.slide_analyzer import SlideAnalyzer
from models.element_type import ElementType
from models.slide_type import SlideType


# ---------------------------------------------------------------------------
# Helpers to build in-memory .pptx files
# ---------------------------------------------------------------------------

def _save_prs(prs: Presentation, tmp_path: Path, name: str = "test.pptx") -> str:
    """Save a python-pptx Presentation to a temp file and return its path."""
    path = tmp_path / name
    prs.save(str(path))
    return str(path)


def _make_title_slide(tmp_path: Path) -> str:
    """One slide with Title + Subtitle placeholders (layout index 0)."""
    prs = Presentation()
    layout = prs.slide_layouts[0]  # 'Title Slide' layout
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "My Presentation"
    # Set subtitle if placeholder exists
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = "Subtitle Text"
    return _save_prs(prs, tmp_path, "title_slide.pptx")


def _make_content_slide(tmp_path: Path) -> str:
    """Two-slide deck: slide 1 = title slide, slide 2 = title+body."""
    prs = Presentation()
    # Slide 1: title slide
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Cover"
    # Slide 2: title + content layout (index 1)
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    for ph in s2.placeholders:
        if ph.placeholder_format.idx == 1:
            tf = ph.text_frame
            tf.text = "First bullet"
            tf.add_paragraph().text = "Second bullet"
    return _save_prs(prs, tmp_path, "content_slide.pptx")


def _make_blank_slide(tmp_path: Path) -> str:
    """One slide using the blank layout (index 6) — no placeholders."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    return _save_prs(prs, tmp_path, "blank_slide.pptx")


def _make_image_only_slide(tmp_path: Path) -> str:
    """One slide with a picture shape added via add_picture, no text."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    # Create a 1×1 white PNG in memory (minimal valid PNG)
    png_bytes = (
        b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01'
        b'\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00'
        b'\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18'
        b'\xd8N\x00\x00\x00\x00IEND\xaeB`\x82'
    )
    img_stream = io.BytesIO(png_bytes)
    slide.shapes.add_picture(img_stream, Inches(1), Inches(1), Inches(3), Inches(3))
    return _save_prs(prs, tmp_path, "image_only.pptx")


def _make_multislide(tmp_path: Path) -> str:
    """Three-slide deck for counting tests."""
    prs = Presentation()
    for layout_idx in [0, 1, 6]:
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        if slide.shapes.title is not None:
            slide.shapes.title.text = f"Slide {layout_idx}"
    return _save_prs(prs, tmp_path, "multi.pptx")


# ---------------------------------------------------------------------------
# Tests: SlideType detection
# ---------------------------------------------------------------------------

class TestSlideTypeDetection:
    """Tests for slide type classification logic."""

    def test_title_slide_detected(self, tmp_path):
        """A slide with Title Slide layout is detected as TITLE_SLIDE."""
        path = _make_title_slide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        assert analyses[0].slide_type == SlideType.TITLE_SLIDE

    def test_content_slide_detected(self, tmp_path):
        """A slide with title + body placeholders is CONTENT_SLIDE."""
        path = _make_content_slide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        # Slide 2 (index 1) has body text
        assert analyses[1].slide_type == SlideType.CONTENT_SLIDE

    def test_blank_slide_detected(self, tmp_path):
        """A slide with no shapes is BLANK_SLIDE."""
        path = _make_blank_slide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        assert analyses[0].slide_type == SlideType.BLANK_SLIDE

    def test_image_only_slide_detected(self, tmp_path):
        """A slide with only a picture and no text is IMAGE_ONLY."""
        path = _make_image_only_slide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        assert analyses[0].slide_type == SlideType.IMAGE_ONLY

    @pytest.mark.skip(
        reason=(
            "Section header requires a layout named 'Section*'. "
            "Programmatic creation is not reliable across default layout sets. "
            "Test this in M6 with fixtures/sample_complex.pptx."
        )
    )
    def test_section_header_detected_by_layout_name(self, tmp_path):
        """A slide whose layout name contains 'Section' is SECTION_HEADER."""


# ---------------------------------------------------------------------------
# Tests: ElementType classification
# ---------------------------------------------------------------------------

class TestElementClassification:
    """Tests that shapes are correctly classified into ElementType values."""

    def test_title_placeholder_classified_as_title(self, tmp_path):
        """Title placeholder → ElementType.TITLE."""
        path = _make_title_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        types = [s.element_type for s in analysis.shapes]
        assert ElementType.TITLE in types

    def test_subtitle_placeholder_classified_as_subtitle(self, tmp_path):
        """Subtitle placeholder → ElementType.SUBTITLE."""
        path = _make_title_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        types = [s.element_type for s in analysis.shapes]
        assert ElementType.SUBTITLE in types

    def test_body_placeholder_classified_as_body(self, tmp_path):
        """Body/content placeholder with text → ElementType.BODY."""
        path = _make_content_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[1]
        types = [s.element_type for s in analysis.shapes]
        assert ElementType.BODY in types

    def test_picture_shape_classified_as_image(self, tmp_path):
        """An add_picture() shape → ElementType.IMAGE."""
        path = _make_image_only_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        types = [s.element_type for s in analysis.shapes]
        assert ElementType.IMAGE in types

    def test_no_grouped_on_simple_slides(self, tmp_path):
        """Simple slides have no GROUPED elements."""
        path = _make_content_slide(tmp_path)
        for analysis in SlideAnalyzer(path).analyze_all():
            types = [s.element_type for s in analysis.shapes]
            assert ElementType.GROUPED not in types


# ---------------------------------------------------------------------------
# Tests: analyze_all() behaviour
# ---------------------------------------------------------------------------

class TestAnalyzeAll:
    """Integration-level tests for the analyze_all() flow."""

    def test_returns_one_analysis_per_slide(self, tmp_path):
        """analyze_all() returns exactly N SlideAnalysis for an N-slide deck."""
        path = _make_multislide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        assert len(analyses) == 3

    def test_slide_numbers_are_one_based(self, tmp_path):
        """SlideAnalysis.slide_number values are 1-based and sequential."""
        path = _make_multislide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        numbers = [a.slide_number for a in analyses]
        assert numbers == [1, 2, 3]

    def test_slide_index_is_zero_based(self, tmp_path):
        """SlideAnalysis.slide_index values are 0-based."""
        path = _make_multislide(tmp_path)
        analyses = SlideAnalyzer(path).analyze_all()
        assert analyses[0].slide_index == 0
        assert analyses[2].slide_index == 2

    def test_total_element_count_matches_shapes_length(self, tmp_path):
        """SlideAnalysis.total_element_count == len(shapes) always."""
        path = _make_content_slide(tmp_path)
        for analysis in SlideAnalyzer(path).analyze_all():
            assert analysis.total_element_count == len(analysis.shapes)

    def test_layout_name_is_captured(self, tmp_path):
        """SlideAnalysis.layout_name is populated when layout has a name."""
        path = _make_title_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        assert analysis.layout_name is not None
        assert isinstance(analysis.layout_name, str)
        assert len(analysis.layout_name) > 0

    def test_invalid_pptx_exits_cleanly(self, tmp_path):
        """A file with a .pptx extension but garbage content → SystemExit."""
        bad = tmp_path / "garbage.pptx"
        bad.write_bytes(b"not a pptx file at all")
        with pytest.raises(SystemExit) as exc_info:
            SlideAnalyzer(str(bad))
        assert exc_info.value.code is not None

    def test_nonexistent_file_exits_cleanly(self):
        """A path to a nonexistent file → SystemExit with clear message."""
        with pytest.raises(SystemExit) as exc_info:
            SlideAnalyzer("/absolutely/nonexistent/path.pptx")
        assert "not found" in str(exc_info.value).lower()

    def test_non_pptx_extension_detected_by_analyzer(self, tmp_path):
        """A file without .pptx extension but opened via SlideAnalyzer → SystemExit."""
        txt = tmp_path / "document.txt"
        txt.write_text("not a pptx")
        with pytest.raises(SystemExit):
            SlideAnalyzer(str(txt))


# ---------------------------------------------------------------------------
# Tests: paragraph counting
# ---------------------------------------------------------------------------

class TestParagraphCounting:
    """Tests for _count_paragraphs accuracy."""

    def test_body_shape_paragraph_count(self, tmp_path):
        """Body placeholder with 2 bullets has paragraph_count == 2."""
        path = _make_content_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[1]
        body_shapes = [s for s in analysis.shapes if s.element_type == ElementType.BODY]
        assert len(body_shapes) >= 1
        assert body_shapes[0].paragraph_count == 2

    def test_title_paragraph_count_is_one(self, tmp_path):
        """Title placeholder with one line has paragraph_count == 1."""
        path = _make_title_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        titles = [s for s in analysis.shapes if s.element_type == ElementType.TITLE]
        assert titles[0].paragraph_count == 1

    def test_image_shape_paragraph_count_is_zero(self, tmp_path):
        """An IMAGE shape has paragraph_count == 0."""
        path = _make_image_only_slide(tmp_path)
        analysis = SlideAnalyzer(path).analyze_all()[0]
        images = [s for s in analysis.shapes if s.element_type == ElementType.IMAGE]
        assert images[0].paragraph_count == 0
